"""生成账号批量导入的测试文件。

不从零拼文件，而是以接口下载的模板为底稿：清空示例行后，套用模板自带
「部门清单 / 角色清单」里的真实取值。这样同一份脚本在本地库与线上库
（两库部门名称可能不同）都能直接跑通，不会因部门名写死而整批校验失败。

用法：python gen_account_import.py <模板路径> <输出路径> <账号A> <账号B>
"""
from __future__ import annotations

import sys

import openpyxl

HEADER_ROW = 4  # 模板结构：说明 3 行 → 表头第 4 行 → 数据第 5 行起
DATA_ROW = 5


def first_value(sheet, col: int = 1) -> str:
    """取清单页某一列的第一个非空值（跳过表头）。"""
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row and len(row) > col and row[col]:
            return str(row[col]).strip()
    return ""


def main() -> int:
    if len(sys.argv) != 5:
        raise SystemExit(__doc__)
    src, dst, acct_a, acct_b = sys.argv[1:5]

    wb = openpyxl.load_workbook(src)
    ws = wb["账号数据"]
    dept = first_value(wb["部门清单"], 1)  # 名称列
    role = first_value(wb["角色清单"], 1)  # 名称列
    if not dept or not role:
        raise SystemExit(f"模板缺少部门或角色清单：dept={dept!r} role={role!r}")

    # 清掉第 5 行起的示例数据，避免跟着一起被导入
    for r in range(DATA_ROW, ws.max_row + 1):
        for c in range(1, 8):
            ws.cell(row=r, column=c).value = None

    rows = [
        ("导入甲", acct_a, "Pass#2026", dept, role, "全公司", "启用"),
        ("导入乙", acct_b, "Pass#2026", dept, role, "本部门", ""),
    ]
    for i, row in enumerate(rows):
        for c, val in enumerate(row, start=1):
            ws.cell(row=DATA_ROW + i, column=c).value = val

    wb.save(dst)
    print(f"[gen] dept={dept} role={role} → {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
