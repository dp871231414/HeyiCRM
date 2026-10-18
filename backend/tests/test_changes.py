"""本轮需求变更的回归测试。

覆盖：
1. 任务管理暴露「创建时间 / 更新时间」且只读；
2. 指标维护范围由权限决定，与角色属性（可见范围）解耦；
3. 统计周期支持范围内任意月份（不限于已有数据的月份）。
"""
from __future__ import annotations

import io
import time
from datetime import datetime

from app.db.session import SessionLocal
from app.models import Task

from tests.conftest import login


def _task_body(dept_id: str = "D01", owner: str = "林雪") -> dict:
    return {
        "parent": "变更测试任务",
        "title": "分类A",
        "description": "",
        "abnormal": False,
        "target": 100,
        "done": 0,
        "unit": "万元",
        "note": "首条进展",
        "start_date": "2026-10-01",
        "end_date": "2026-12-31",
        "dept_id": dept_id,
        "owner": owner,
    }


# ---------------------------------------------------------------- 任务时间字段


def test_task_list_exposes_timestamps(client, admin_headers):
    """列表返回 createdAt / updatedAt，供前端只读展示。"""
    items = client.get("/api/v1/tasks", headers=admin_headers).json()["items"]
    assert items, "种子数据应含任务"
    for t in items:
        assert "createdAt" in t and "updatedAt" in t
        assert t["createdAt"], "创建时间不应为空（系统自动记录）"


def test_task_update_touches_updated_at(client, admin_headers):
    created = client.post("/api/v1/tasks", json=_task_body(), headers=admin_headers).json()
    before = created["updatedAt"]
    assert before

    body = _task_body()
    body["done"] = 60
    updated = client.patch(f"/api/v1/tasks/{created['id']}", json=body, headers=admin_headers).json()
    assert updated["done"] == 60
    assert updated["updatedAt"] >= before
    # 创建时间不被修改
    assert updated["createdAt"] == created["createdAt"]

    client.delete(f"/api/v1/tasks/{created['id']}", headers=admin_headers)


def test_task_timestamps_not_writable_by_client(client, admin_headers):
    """请求体里塞时间字段应被忽略（只读，由系统记录）。"""
    body = _task_body()
    body["createdAt"] = "1999-01-01T00:00:00+00:00"
    body["updatedAt"] = "1999-01-01T00:00:00+00:00"
    created = client.post("/api/v1/tasks", json=body, headers=admin_headers).json()
    assert not created["createdAt"].startswith("1999")
    client.delete(f"/api/v1/tasks/{created['id']}", headers=admin_headers)


# ---------------------------------------------------------------- 权限可配置


def test_permissions_catalog_exposes_new_metric_keys(client, admin_headers):
    groups = client.get("/api/v1/permissions", headers=admin_headers).json()["groups"]
    metric = next(g for g in groups if g["menu"] == "metric")
    keys = {i["key"] for i in metric["items"]}
    assert {"metric.edit", "metric.edit_company", "metric.edit_project"} <= keys


def test_board_reports_edit_perms(client, admin_headers):
    d = client.get("/api/v1/metrics/board", params={"org_id": "CO"}, headers=admin_headers).json()
    assert d["perms"] == {"edit": True, "editCompany": True, "editProject": True}


def test_company_edit_gated_by_permission_not_scope(client, admin_headers):
    """仅去掉 metric.edit_company 后，即使可见范围仍是「全公司」也不能维护公司整体目标。

    这条断言正是本次改造的核心：操作门槛只看权限，不看角色属性。
    """
    from app.db.session import SessionLocal
    from app.models import Role

    with SessionLocal() as db:
        role = db.get(Role, "admin")
        original = list(role.permissions or [])
        scope_before = role.scope
        role.permissions = [p for p in original if p != "metric.edit_company"]
        db.commit()
        try:
            d = client.get("/api/v1/metrics/board", params={"org_id": "CO"}, headers=admin_headers).json()
            assert d["perms"]["edit"] is True
            assert d["perms"]["editCompany"] is False
            # 可见范围未变（仍为全公司），但编辑权已被权限收回
            assert d["canEditAll"] is False
        finally:
            role.permissions = original
            role.scope = scope_before
            db.commit()


def test_project_edit_gated_by_permission(client, admin_headers):
    from app.db.session import SessionLocal
    from app.models import Role

    with SessionLocal() as db:
        role = db.get(Role, "admin")
        original = list(role.permissions or [])
        role.permissions = [p for p in original if p != "metric.edit_project"]
        db.commit()
        try:
            d = client.get("/api/v1/metrics/board", params={"org_id": "CO"}, headers=admin_headers).json()
            assert d["editableProjects"] == []
        finally:
            role.permissions = original
            db.commit()


def test_dept_head_cannot_edit_company_targets(client, head_headers):
    """部门负责人无 metric.edit_company，公司整体目标只读。"""
    d = client.get("/api/v1/metrics/board", params={"org_id": "D06"}, headers=head_headers).json()
    assert d["canEditAll"] is False
    assert d["perms"]["editCompany"] is False


def test_role_is_system_is_configurable(client, admin_headers):
    """内置标记可配置：新建角色可设为内置，且可再取消。"""
    payload = {
        "name": "可配置内置角色-测试",
        "description": "",
        "scope": "self",
        "permissions": ["dashboard.view"],
        "isSystem": True,
    }
    created = client.post("/api/v1/roles", json=payload, headers=admin_headers).json()
    assert created["isSystem"] is True

    # 内置状态下不可删除
    blocked = client.delete(f"/api/v1/roles/{created['id']}", headers=admin_headers)
    assert blocked.status_code == 409

    # 取消内置后即可删除
    off = dict(payload, name="可配置内置角色-测试", isSystem=False)
    assert client.patch(f"/api/v1/roles/{created['id']}", json=off, headers=admin_headers).json()["isSystem"] is False
    assert client.delete(f"/api/v1/roles/{created['id']}", headers=admin_headers).status_code == 200


def test_role_update_without_is_system_keeps_value(client, admin_headers):
    """未携带 isSystem 时保持原值，避免被误清空。"""
    roles = client.get("/api/v1/roles", headers=admin_headers).json()
    admin = next(r for r in roles if r["id"] == "admin")
    assert admin["isSystem"] is True

    payload = {
        "name": admin["name"],
        "description": admin["description"],
        "scope": admin["scope"],
        "permissions": admin["permissions"],
    }
    updated = client.patch("/api/v1/roles/admin", json=payload, headers=admin_headers).json()
    assert updated["isSystem"] is True


# ---------------------------------------------------------------- 周期范围


def test_board_accepts_month_without_data(client, admin_headers):
    """统计周期可选范围内任意月份，无数据也应正常返回而非报错。"""
    for period in ("2026-09", "2099-12"):
        resp = client.get(
            "/api/v1/metrics/board", params={"org_id": "CO", "period": period}, headers=admin_headers
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["period"] == period


# ---------------------------------------------------------------- 模板与导出（本轮新增）

_MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _load_xlsx(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content))


def _header_of(ws, n: int, row: int) -> list:
    return [ws.cell(row=row, column=i).value for i in range(1, n + 1)]


def _export_file_bytes(client, headers, export_url: str, object_type: str) -> bytes:
    """触发导出（仅生成下载中心记录），取回 xlsx 字节并清理该记录。

    回归：导出端点不再即时返回文件，而是返回 DownloadRecordOut，
    实际文件需到「下载中心」按 id 下载。
    """
    rec = client.get(export_url, headers=headers).json()
    assert rec["objectType"] == object_type, f"导出应返回下载记录，实际：{rec}"
    rid = rec["id"]
    try:
        resp = client.get(f"/api/v1/downloads/{rid}/file", headers=headers)
        assert resp.status_code == 200, resp.text
        return resp.content
    finally:
        client.delete(f"/api/v1/downloads/{rid}", headers=headers)


def test_task_template_columns_match_list_and_form(client, admin_headers):
    """模板列名与「任务管理」列表/表单同一套叫法，且补齐了「数量单位」。"""
    from app.services.importer import HEADER_ROW, TASK_COLUMNS

    resp = client.get("/api/v1/tasks/import/template", headers=admin_headers)
    assert resp.status_code == 200
    ws = _load_xlsx(resp.content).worksheets[0]
    assert _header_of(ws, len(TASK_COLUMNS), HEADER_ROW) == TASK_COLUMNS
    # 不再使用「一级任务 / 二级任务」这类仅模板可见的旧叫法
    assert "一级任务" not in TASK_COLUMNS and "二级任务" not in TASK_COLUMNS
    # 表单里有、原来模板缺失的字段
    for name in ("任务名称", "任务分类", "任务描述", "数量单位"):
        assert name in TASK_COLUMNS


def test_task_template_styling(client, admin_headers):
    """说明栏 + 表头背景色 + 冻结表头。"""
    from app.services.importer import _BLUE_MAIN, HEADER_ROW, TASK_COLUMNS

    ws = _load_xlsx(client.get("/api/v1/tasks/import/template", headers=admin_headers).content).worksheets[0]
    banner = ws.cell(row=1, column=1).value
    assert banner and "模板" in str(banner), "顶部应有说明栏标题"
    head = ws.cell(row=HEADER_ROW, column=1)
    assert str(head.fill.fgColor.rgb).endswith(_BLUE_MAIN), "字段名称栏应有背景色"
    assert head.font.bold and ws.freeze_panes == f"A{HEADER_ROW + 1}"


def test_task_template_notes_mention_unit_and_system_columns(client, admin_headers):
    wb = _load_xlsx(client.get("/api/v1/tasks/import/template", headers=admin_headers).content)
    notes = wb["填写说明"]
    txt = " ".join(str(notes.cell(row=r, column=1).value or "") for r in range(1, notes.max_row + 1))
    assert "数量单位" in txt
    assert "创建人" in txt and "创建时间" in txt


def test_task_export_columns_align_with_list(client, admin_headers):
    """导出列 = 任务ID + 全部可填写列 + 系统列；业务列与导入模板同名。"""
    from app.services.importer import HEADER_ROW, TASK_COLUMNS, TASK_EXPORT_COLUMNS

    ws = _load_xlsx(_export_file_bytes(client, admin_headers, "/api/v1/tasks/export", "task")).worksheets[0]
    header = _header_of(ws, len(TASK_EXPORT_COLUMNS), HEADER_ROW)
    assert header == TASK_EXPORT_COLUMNS
    for name in TASK_COLUMNS:
        assert name in header
    for name in ("任务ID", "完成进度", "状态", "创建人", "创建时间", "更新时间"):
        assert name in header


def test_task_export_drops_dept_column_and_keeps_creator(client, admin_headers):
    """任务导出不再单列「责任部门」（部门由责任人自动带出），但「创建人」仍为姓名。"""
    from app.services.importer import HEADER_ROW, TASK_EXPORT_COLUMNS

    assert "责任部门" not in TASK_EXPORT_COLUMNS, "责任部门已改为由责任人自动带出，导出模板不应再单列该列"
    ws = _load_xlsx(_export_file_bytes(client, admin_headers, "/api/v1/tasks/export", "task")).worksheets[0]
    idx = {name: i + 1 for i, name in enumerate(TASK_EXPORT_COLUMNS)}
    creators = {
        ws.cell(row=r, column=idx["创建人"]).value
        for r in range(HEADER_ROW + 1, ws.max_row + 1)
    }
    assert any(c for c in creators), "创建人应有姓名"


def test_task_import_derives_dept_from_owner(client, admin_headers):
    """导入仅填责任人（不填责任部门）时，责任部门应自动带出为该责任人的所属部门。"""
    from openpyxl import Workbook

    from app.services.importer import TASK_COLUMNS

    # 取责任人及其所属部门（联动账号管理）
    owners = {o["name"]: o for o in client.get("/api/v1/meta/owners", headers=admin_headers).json()}
    owner = owners["郑凯"]

    wb = Workbook()
    ws = wb.active
    ws.append(TASK_COLUMNS)
    vals = {
        "任务名称": "导入派生部门任务",
        "任务分类": "派生子组",
        "任务描述": "描述",
        "任务进展": "首条进展",
        "是否异常": "否",
        "目标": 100,
        "已完成": 50,
        "计划开始日期": "2026-01-01",
        "计划结束日期": "2026-12-31",
        "责任人": "郑凯",
    }
    ws.append([vals.get(c, "") for c in TASK_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        "/api/v1/tasks/import",
        files={"file": ("t.xlsx", buf.getvalue(), _MIME_XLSX)},
        headers=admin_headers,
    )
    assert resp.status_code == 200 and resp.json()["imported"] == 1
    items = client.get("/api/v1/tasks", params={"keyword": "导入派生部门任务"}, headers=admin_headers).json()["items"]
    assert items, "应导入成功"
    try:
        assert items[0]["deptId"] == owner["deptId"], "责任部门应自动等于责任人的所属部门"
        assert items[0]["deptName"] == owner["deptName"]
    finally:
        for it in items:
            client.delete(f"/api/v1/tasks/{it['id']}", headers=admin_headers)


def test_task_export_is_reimportable(client, admin_headers):
    """导出文件可直接当导入模板回灌（列齐全，能通过解析）。"""
    from app.services.importer import TASK_COLUMNS, TASK_OPTIONAL_COLUMNS, parse_rows

    content = _export_file_bytes(client, admin_headers, "/api/v1/tasks/export", "task")
    rows = parse_rows(content, TASK_COLUMNS, 2000, optional=TASK_OPTIONAL_COLUMNS)
    assert rows, "导出内容应能被解析出数据行"
    assert all(r["任务名称"] for r in rows)


def test_parse_rows_skips_banner_row():
    """模板顶部有说明栏时，表头行应被自动定位。"""
    from openpyxl import Workbook

    from app.services.importer import TASK_COLUMNS, TASK_OPTIONAL_COLUMNS, parse_rows

    wb = Workbook()
    ws = wb.active
    ws.append(["重点任务导入标准模板"])
    ws.append(["请勿修改第 4 行字段名称"])
    ws.append([])
    ws.append(TASK_COLUMNS)
    vals = {
        "任务名称": "带说明栏的任务",
        "目标": 1,
        "已完成": 0,
        "计划开始日期": "2026-01-01",
        "计划结束日期": "2026-12-31",
        "责任部门": "市场部",
        "责任人": "郑凯",
    }
    ws.append([vals.get(c, "") for c in TASK_COLUMNS])
    buf = io.BytesIO()
    wb.save(buf)

    rows = parse_rows(buf.getvalue(), TASK_COLUMNS, 100, optional=TASK_OPTIONAL_COLUMNS)
    assert len(rows) == 1
    assert rows[0]["任务名称"] == "带说明栏的任务"


def test_parse_rows_tolerates_legacy_template_missing_optional_column():
    """老模板缺少「数量单位」等可选列时仍可导入，按空值处理。"""
    from openpyxl import Workbook

    from app.services.importer import TASK_COLUMNS, TASK_OPTIONAL_COLUMNS, parse_rows

    legacy = [c for c in TASK_COLUMNS if c != "数量单位"]
    wb = Workbook()
    ws = wb.active
    ws.append(legacy)
    vals = {
        "任务名称": "老模板任务",
        "目标": 1,
        "已完成": 0,
        "计划开始日期": "2026-01-01",
        "计划结束日期": "2026-12-31",
        "责任部门": "市场部",
        "责任人": "郑凯",
    }
    ws.append([vals.get(c, "") for c in legacy])
    buf = io.BytesIO()
    wb.save(buf)

    rows = parse_rows(buf.getvalue(), TASK_COLUMNS, 100, optional=TASK_OPTIONAL_COLUMNS)
    assert rows[0]["任务名称"] == "老模板任务"
    assert rows[0].get("数量单位", "") == ""


def test_parse_rows_still_rejects_missing_required_column():
    """必填列缺失仍需报错。"""
    from openpyxl import Workbook

    from app.services.importer import TASK_COLUMNS, TASK_OPTIONAL_COLUMNS, parse_rows

    cols = ["任务名称", "目标"]  # 缺少日期、部门、责任人等必填列
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    ws.append(["x", 1])
    buf = io.BytesIO()
    wb.save(buf)

    try:
        parse_rows(buf.getvalue(), TASK_COLUMNS, 100, optional=TASK_OPTIONAL_COLUMNS)
    except ValueError as exc:
        assert "模板缺少列" in str(exc)
    else:
        raise AssertionError("应抛出缺少列的 ValueError")


def test_task_list_exposes_creator_name(client, admin_headers):
    """列表返回创建人姓名，供新增的「创建人」列展示。"""
    items = client.get("/api/v1/tasks", headers=admin_headers).json()["items"]
    assert items
    for t in items:
        assert "createdById" in t and "createdByName" in t
    assert any(t["createdByName"] for t in items)


def test_template_has_no_blank_row_after_header(client, admin_headers):
    """表头下一行必须直接是示例行。

    回归：freeze_panes 若用 ws.cell(...) 会创建空单元格，导致数据整体下移一行、
    模板里多出一行空白（解析时虽能跳过，但肉眼可见地不整洁）。
    """
    from app.services.importer import HEADER_ROW, TASK_COLUMNS

    ws = _load_xlsx(client.get("/api/v1/tasks/import/template", headers=admin_headers).content).worksheets[0]
    assert ws.cell(row=HEADER_ROW, column=1).value == TASK_COLUMNS[0]
    assert ws.cell(row=HEADER_ROW + 1, column=1).value, "表头下一行应直接是示例行"
    assert ws.freeze_panes == f"A{HEADER_ROW + 1}"


def test_export_has_no_blank_row_after_header(client, admin_headers):
    from app.services.importer import HEADER_ROW

    ws = _load_xlsx(_export_file_bytes(client, admin_headers, "/api/v1/tasks/export", "task")).worksheets[0]
    assert ws.cell(row=HEADER_ROW, column=1).value == "任务ID"
    assert ws.cell(row=HEADER_ROW + 1, column=1).value, "表头下一行应直接是数据"


# ---------------------------------------------------------------- 项目公司权限与删除（本轮新增）


def test_perm_catalog_exposes_project_actions(client, admin_headers):
    """权限目录暴露项目公司的增 / 改 / 删三个独立权限项，供角色按需勾选。"""
    catalog = client.get("/api/v1/permissions", headers=admin_headers).json()
    keys = {item["key"] for g in catalog["groups"] for item in g["items"]}
    assert {"project.create", "project.edit", "project.delete"} <= keys


def test_admin_role_includes_project_actions(client, admin_headers):
    roles = {r["id"]: r for r in client.get("/api/v1/roles", headers=admin_headers).json()}
    assert {"project.create", "project.edit", "project.delete"} <= set(roles["admin"]["permissions"])


def test_role_without_project_perm_cannot_manage_projects(client, admin_headers, head_headers):
    """部门负责人可维护项目公司指标（metric.edit_project），但没有项目公司管理权限。

    回归点：此前增删改的门是「具备公司整体目标权限且当前视角为全公司」，
    导致这类角色看得到数据却无法操作。现改为纯权限点判定。
    """
    denied_create = client.post(
        "/api/v1/projects", json={"name": "不应创建", "head": "x", "targets": []}, headers=head_headers
    )
    assert denied_create.status_code == 403
    assert client.patch("/api/v1/projects/PC1", json={"name": "不应改名", "head": "x"}, headers=head_headers).status_code == 403
    assert client.delete("/api/v1/projects/PC1", headers=head_headers).status_code == 403
    # 未被波及
    assert any(p["id"] == "PC1" for p in client.get("/api/v1/projects", headers=admin_headers).json())


def test_project_manage_configurable_via_custom_role(client, admin_headers):
    """给自定义角色勾上项目公司权限后即可操作——能力由权限决定，与角色标识无关。"""
    import time

    suffix = str(int(time.time() * 1000))[-6:]
    role = client.post(
        "/api/v1/roles",
        json={
            "name": f"项目管理员{suffix}",
            "description": "",
            "scope": "all",
            "permissions": ["metric.view", "project.create", "project.edit", "project.delete"],
        },
        headers=admin_headers,
    )
    assert role.status_code == 201, role.text
    rid = role.json()["id"]

    account = f"pm{suffix}"
    created_user = client.post(
        "/api/v1/users",
        json={
            "name": f"项目管理员{suffix}",
            "account": account,
            "password": "123456",
            "dept_id": "CO",
            "role_id": rid,
            "scope": "all",
            "status": "active",
        },
        headers=admin_headers,
    )
    assert created_user.status_code == 201, created_user.text
    uid = created_user.json()["id"]

    try:
        headers, _ = login(client, account)
        made = client.post(
            "/api/v1/projects", json={"name": f"权限验证公司{suffix}", "head": "x", "targets": []}, headers=headers
        )
        assert made.status_code == 201, made.text
        pid = made.json()["data"]["id"]

        renamed = client.patch(f"/api/v1/projects/{pid}", json={"name": f"权限验证公司{suffix}b", "head": "y"}, headers=headers)
        assert renamed.status_code == 200, renamed.text
        assert client.delete(f"/api/v1/projects/{pid}", headers=headers).status_code == 200
    finally:
        client.delete(f"/api/v1/users/{uid}", headers=admin_headers)
        client.delete(f"/api/v1/roles/{rid}", headers=admin_headers)


def test_project_delete_reassigns_dangling_refs(client, admin_headers):
    """删除项目公司：关联账号与子组织迁移到上级组织，不留悬空引用。

    回归点：此前只把账号置为停用、未改 dept_id，删除后 dept_id 会指向不存在的组织。
    """
    from sqlalchemy import select

    from app.db.session import SessionLocal
    from app.models import Organization, User

    made = client.post(
        "/api/v1/projects", json={"name": "待删除项目公司X", "head": "测试", "targets": []}, headers=admin_headers
    )
    assert made.status_code == 201, made.text
    pid = made.json()["data"]["id"]

    with SessionLocal() as db:
        db.add(Organization(id="DPX", name="待删除项目公司的子部门", type="dept", parent_id=pid, head="—", sort_order=900))
        db.add(
            User(
                id="U9X",
                name="临时账号X",
                account="tmpx",
                password_hash="x",
                dept_id=pid,
                role_id="member",
                status="active",
            )
        )
        db.commit()

    assert client.delete(f"/api/v1/projects/{pid}", headers=admin_headers).status_code == 200

    with SessionLocal() as db:
        assert db.get(Organization, pid) is None, "项目公司应已删除"
        child = db.get(Organization, "DPX")
        moved = db.get(User, "U9X")
        assert child is not None and child.parent_id == "CO", "子组织应上移到 CO，而非悬空"
        assert moved is not None and moved.dept_id == "CO", "账号归属应迁移到 CO"
        assert moved.status == "disabled", "账号应被停用而非删除"

        # 全库不应存在指向已删除组织的引用
        org_ids = {o.id for o in db.scalars(select(Organization)).all()}
        assert all(u.dept_id in org_ids for u in db.scalars(select(User)).all())
        assert all(o.parent_id is None or o.parent_id in org_ids for o in db.scalars(select(Organization)).all())

        # 清理临时数据
        if child is not None:
            db.delete(child)
        if moved is not None:
            db.delete(moved)
        db.commit()


# ---------------------------------------------------------------- 任务分页与状态过滤（本轮新增）


def test_task_status_filter_matches_computed_status(client, admin_headers):
    """status 过滤下推到 SQL 后，结果必须与 Python 侧判定的 statusLevel 完全一致。

    回归点：分页在数据库层、过滤在 Python 层时，页数与总数会对不上；
    下推后两者必须仍语义一致。
    """
    all_items = client.get("/api/v1/tasks", params={"size": 2000}, headers=admin_headers).json()
    by_level: dict[str, set[str]] = {}
    for it in all_items["items"]:
        by_level.setdefault(it["statusLevel"], set()).add(it["id"])

    for level in ("done", "normal", "due_soon", "overdue", "risk"):
        got = client.get("/api/v1/tasks", params={"status": level, "size": 2000}, headers=admin_headers).json()
        got_ids = {it["id"] for it in got["items"]}
        assert got_ids == by_level.get(level, set()), f"{level} 的 SQL 过滤结果与计算状态不一致"
        assert got["total"] == len(got_ids), f"{level} 的 total 应为过滤后的总数"
        assert all(it["statusLevel"] == level for it in got["items"])


def test_task_list_pagination_is_at_sql_level(client, admin_headers):
    """分页在 SQL 层生效：单页条数受限、翻页不重不漏、越界页为空。"""
    total = client.get("/api/v1/tasks", params={"size": 1}, headers=admin_headers).json()["total"]
    assert total >= 3, "演示数据任务数不足，无法验证翻页"

    p1 = client.get("/api/v1/tasks", params={"page": 1, "size": 2}, headers=admin_headers).json()
    p2 = client.get("/api/v1/tasks", params={"page": 2, "size": 2}, headers=admin_headers).json()
    assert len(p1["items"]) == 2
    assert p1["total"] == total and p2["total"] == total
    assert p1["page"] == 1 and p2["page"] == 2
    assert not ({i["id"] for i in p1["items"]} & {i["id"] for i in p2["items"]}), "不同页不应出现重复项"

    # 页序稳定：逐页取回拼起来应与一次性取回的顺序一致
    all_ids = [i["id"] for i in client.get("/api/v1/tasks", params={"size": 2000}, headers=admin_headers).json()["items"]]
    paged_ids = [i["id"] for i in p1["items"]] + [i["id"] for i in p2["items"]]
    assert paged_ids == all_ids[: len(paged_ids)], "分页顺序应与整体顺序一致"

    empty = client.get("/api/v1/tasks", params={"page": 999, "size": 2}, headers=admin_headers).json()
    assert empty["items"] == [] and empty["total"] == total, "越界页应返回空列表但总数不变"


def test_task_list_page_size_has_upper_bound(client, admin_headers):
    """size 有上限护栏，避免一次拉取过多数据。"""
    assert client.get("/api/v1/tasks", params={"size": 2000}, headers=admin_headers).status_code == 200
    assert client.get("/api/v1/tasks", params={"size": 99999}, headers=admin_headers).status_code == 422
    assert client.get("/api/v1/tasks", params={"size": 0}, headers=admin_headers).status_code == 422


def test_task_list_filters_can_combine_with_pagination(client, admin_headers):
    """状态过滤 + 分页组合时，total 是过滤后的总数（而不是全表总数）。"""
    all_items = client.get("/api/v1/tasks", params={"size": 2000}, headers=admin_headers).json()["items"]
    dept_id = all_items[0]["deptId"]
    in_dept = [i for i in all_items if i["deptId"] == dept_id]

    got = client.get("/api/v1/tasks", params={"dept_id": dept_id, "size": 1}, headers=admin_headers).json()
    assert got["total"] == len(in_dept), "total 应为该部门的任务数"
    assert len(got["items"]) == 1, "单页条数应受 size 限制"


# ------------------------------------------------- 指标编辑权限分块 + 批量趋势（本轮新增）


def _role_user(client, admin_headers, perms, tag):
    """建一个角色 + 一个挂在 CO 的账号，返回 (role_id, user_id, account)。"""
    import time

    suffix = f"{tag}{str(int(time.time() * 1000))[-5:]}"
    role = client.post(
        "/api/v1/roles",
        json={"name": f"指标权限{suffix}", "description": "", "scope": "all", "permissions": perms},
        headers=admin_headers,
    )
    assert role.status_code == 201, role.text
    rid = role.json()["id"]
    created = client.post(
        "/api/v1/users",
        json={
            "name": f"指标权限{suffix}",
            "account": f"mp{suffix}",
            "password": "123456",
            "dept_id": "CO",
            "role_id": rid,
            "scope": "all",
            "status": "active",
        },
        headers=admin_headers,
    )
    assert created.status_code == 201, created.text
    return rid, created.json()["id"], f"mp{suffix}"


def test_perm_catalog_groups_metric_blocks(client, admin_headers):
    """权限目录里，指标维护的编辑权限按页面三块分组呈现。"""
    catalog = client.get("/api/v1/permissions", headers=admin_headers).json()
    metric = next(g for g in catalog["groups"] if g["menu"] == "metric")
    grouped: dict[str, list[str]] = {}
    for it in metric["items"]:
        if it.get("group"):
            grouped.setdefault(it["group"], []).append(it["key"])

    assert set(grouped) == {"公司整体目标", "指标监控", "项目公司指标"}
    assert grouped["公司整体目标"] == ["metric.edit_company"]
    assert grouped["指标监控"] == ["metric.edit"]
    assert set(grouped["项目公司指标"]) == {
        "metric.edit_project",
        "project.create",
        "project.edit",
        "project.delete",
    }


def test_metric_edits_are_independent_per_block(client, admin_headers):
    """三块编辑权限互相独立，没有隐含的全局前置权限。

    回归点：此前 metric.edit 是所有指标编辑的前置条件 —— 即使勾了
    「维护公司整体目标」，缺 metric.edit 也无法编辑，页面上看不出这层依赖。
    """
    period = "2026-10"

    # 只给「维护公司整体目标」
    rid, uid, account = _role_user(client, admin_headers, ["metric.view", "metric.edit_company"], "co")
    try:
        headers, _ = login(client, account)
        # 公司整体目标（annual 指标 + 公司组织）→ 应允许
        ok = client.put(
            "/api/v1/metrics/targets",
            json={"metric_id": "REV_TRANS", "org_id": "CO", "period": period, "value": 1234},
            headers=headers,
        )
        assert ok.status_code == 200, ok.text

        # 指标监控 → 应拒绝，且提示缺的是哪个权限
        denied_monitor = client.put(
            "/api/v1/metrics/actuals",
            json={"metric_id": "INC_TOT", "org_id": "CO", "period": period, "value": 3},
            headers=headers,
        )
        assert denied_monitor.status_code == 403, denied_monitor.text
        assert "监控指标" in denied_monitor.json()["detail"]

        # 项目公司指标 → 应拒绝
        denied_project = client.put(
            "/api/v1/metrics/targets",
            json={"metric_id": "OP_REV", "org_id": "PC1", "period": period, "value": 100},
            headers=headers,
        )
        assert denied_project.status_code == 403, denied_project.text
    finally:
        client.delete(f"/api/v1/users/{uid}", headers=admin_headers)
        client.delete(f"/api/v1/roles/{rid}", headers=admin_headers)

    # 只给「维护监控指标」
    rid2, uid2, account2 = _role_user(client, admin_headers, ["metric.view", "metric.edit"], "mo")
    try:
        headers2, _ = login(client, account2)
        ok2 = client.put(
            "/api/v1/metrics/actuals",
            json={"metric_id": "INC_TOT", "org_id": "CO", "period": period, "value": 5},
            headers=headers2,
        )
        assert ok2.status_code == 200, ok2.text

        denied_annual = client.put(
            "/api/v1/metrics/targets",
            json={"metric_id": "REV_TRANS", "org_id": "CO", "period": period, "value": 999},
            headers=headers2,
        )
        assert denied_annual.status_code == 403, denied_annual.text
        assert "公司整体目标" in denied_annual.json()["detail"]
    finally:
        client.delete(f"/api/v1/users/{uid2}", headers=admin_headers)
        client.delete(f"/api/v1/roles/{rid2}", headers=admin_headers)


def test_trends_batch_matches_single(client, admin_headers):
    """批量趋势接口的结果必须与逐个调用完全一致，且保持请求顺序。"""
    ids = ["INC_TOT", "FAU_TOT", "RD_CNT", "INC_AMT"]
    singles = {
        mid: client.get(
            f"/api/v1/dashboard/metrics/{mid}/trend", params={"org_id": "CO"}, headers=admin_headers
        ).json()
        for mid in ids
    }
    batch = client.get(
        "/api/v1/dashboard/metrics/trends", params={"ids": ",".join(ids), "org_id": "CO"}, headers=admin_headers
    ).json()

    assert [s["metricId"] for s in batch["series"]] == ids, "应保持调用方给定的顺序"
    for s in batch["series"]:
        assert s == singles[s["metricId"]], f"{s['metricId']} 的批量结果与单个接口不一致"
        assert s["points"], f"{s['metricId']} 应有趋势点"


def test_trends_batch_edge_cases(client, admin_headers):
    """空 ids 返回空列表；不存在的指标被跳过而不是报错。"""
    empty = client.get("/api/v1/dashboard/metrics/trends", params={"ids": ""}, headers=admin_headers).json()
    assert empty == {"series": []}

    partial = client.get(
        "/api/v1/dashboard/metrics/trends",
        params={"ids": "NOT_EXIST,INC_TOT", "org_id": "CO"},
        headers=admin_headers,
    ).json()
    assert [s["metricId"] for s in partial["series"]] == ["INC_TOT"]


# ------------------------------------------------- 组织编辑：返回真实数据 + 公司节点可编辑


def test_org_update_returns_real_payload(client, admin_headers):
    """更新接口要返回真实的 depth 与统计数。

    前端保存后用它就地替换列表节点（不再重拉整棵树），
    若返回占位值（原先 update 恒为 depth=0、各项计数为 0），
    列表的缩进层级与统计数字会立即错乱。
    """
    orgs = client.get("/api/v1/orgs", params={"flat": 1}, headers=admin_headers).json()
    dept = next(o for o in orgs if o["type"] == "dept")

    resp = client.patch(
        f"/api/v1/orgs/{dept['id']}",
        json={"name": dept["name"], "head": "测试负责人"},
        headers=admin_headers,
    )
    assert resp.status_code == 200, resp.text
    body = resp.json()

    assert body["depth"] == dept["depth"], "depth 应与组织树中的层级一致"
    assert body["userCount"] == dept["userCount"]
    assert body["childCount"] == dept["childCount"]
    assert body["taskCount"] == dept["taskCount"]
    assert body["head"] == "测试负责人"
    assert body["parentId"] == dept["parentId"]


def test_company_node_is_editable_but_not_deletable(client, admin_headers):
    """「全公司」根节点可以改名称与负责人，但不能删除。"""
    orgs = client.get("/api/v1/orgs", params={"flat": 1}, headers=admin_headers).json()
    co = next(o for o in orgs if o["id"] == "CO")

    edited = client.patch(
        "/api/v1/orgs/CO", json={"name": co["name"], "head": "总经理"}, headers=admin_headers
    )
    assert edited.status_code == 200, edited.text
    body = edited.json()
    assert body["id"] == "CO"
    assert body["depth"] == 0, "根节点层级应为 0"
    assert body["head"] == "总经理"

    assert client.delete("/api/v1/orgs/CO", headers=admin_headers).status_code == 409, "根节点不可删除"


def test_create_org_returns_real_depth(client, admin_headers):
    """新建子部门返回的 depth 应为真实层级（原先硬编码为 1）。"""
    import time

    orgs = client.get("/api/v1/orgs", params={"flat": 1}, headers=admin_headers).json()
    parent = next(o for o in orgs if o["type"] == "dept")
    name = f"临时子部门{str(int(time.time() * 1000))[-5:]}"

    created = client.post(
        "/api/v1/orgs",
        json={"name": name, "parent_id": parent["id"], "head": "临时"},
        headers=admin_headers,
    )
    assert created.status_code == 201, created.text
    body = created.json()
    try:
        assert body["depth"] == parent["depth"] + 1, "子部门层级应为父级 + 1（而不是固定的 1）"
        assert body["parentId"] == parent["id"]
        assert body["childCount"] == 0 and body["userCount"] == 0
    finally:
        client.delete(f"/api/v1/orgs/{body['id']}", headers=admin_headers)


# ------------------------------ 数据权限边界：仅任务模块按范围隔离（本轮口径）


def test_dashboard_open_by_perm_not_scope(client, head_headers):
    """经营驾驶舱只按角色权限开放：跨组织不拦截，也不按可见范围裁剪。

    回归点：原先部门负责人访问全公司口径的总结 / 总览 / 趋势会被拦成 403。
    现按既定口径放开 —— 数据范围（scope）仅作用于任务模块。
    """
    for path in ("/api/v1/dashboard/summary", "/api/v1/dashboard/overview"):
        r = client.get(path, params={"org_id": "CO"}, headers=head_headers)
        assert r.status_code == 200, f"{path} 不应按数据范围拦截"

    assert (
        client.get("/api/v1/dashboard/metrics/INC_TOT/trend", params={"org_id": "CO"}, headers=head_headers).status_code
        == 200
    )
    assert (
        client.get(
            "/api/v1/dashboard/metrics/trends", params={"ids": "INC_TOT", "org_id": "CO"}, headers=head_headers
        ).status_code
        == 200
    )


def test_dashboard_task_stats_are_global(client, admin_headers, head_headers):
    """驾驶舱里的任务统计按全量口径（与管理员一致），不随可见范围收窄。"""
    admin = client.get("/api/v1/dashboard/dept-task-stats", headers=admin_headers).json()
    head = client.get("/api/v1/dashboard/dept-task-stats", headers=head_headers).json()
    assert head == admin, "驾驶舱任务统计应为全量口径"

    overview_admin = client.get("/api/v1/dashboard/overview", params={"org_id": "CO"}, headers=admin_headers).json()
    overview_head = client.get("/api/v1/dashboard/overview", params={"org_id": "CO"}, headers=head_headers).json()
    assert overview_head["counts"] == overview_admin["counts"]


def test_task_module_still_isolates_by_scope(client, admin_headers, head_headers):
    """任务模块是本系统唯一按数据可见范围隔离的模块（本轮口径的边界）。"""
    admin_items = client.get("/api/v1/tasks", params={"size": 2000}, headers=admin_headers).json()["items"]
    head_items = client.get("/api/v1/tasks", params={"size": 2000}, headers=head_headers).json()["items"]
    admin_ids = {t["id"] for t in admin_items}
    head_ids = {t["id"] for t in head_items}

    assert head_ids < admin_ids, "任务列表仍应按可见范围收窄"
    hidden = sorted(admin_ids - head_ids)
    assert hidden, "应存在范围外不可见的任务"
    # 详情同样受范围约束
    assert client.get(f"/api/v1/tasks/{hidden[0]}", headers=head_headers).status_code == 403


def test_projects_list_not_scope_filtered(client, head_headers):
    """项目公司列表按全公司口径返回，不按可见范围裁剪。"""
    ids = {p["id"] for p in client.get("/api/v1/projects", headers=head_headers).json()}
    assert {"PC1", "PC2", "PC3"} <= ids


# ---------------------- 进度状态口径统一：驾驶舱与指标维护必须给出同一状态 ----------------------


def test_progress_tone_bands():
    """progress_tone 的分档规则（全系统唯一状态口径）。"""
    from app.services.achievement import progress_tone

    tp = 0.8
    assert progress_tone(None, tp) == "none", "无达成率 → 未录入"
    assert progress_tone(1.0, tp) == "done", "达成率 100% → 已完成"
    assert progress_tone(1.35, tp) == "done", "超额完成仍是已完成（优先于比值）"
    assert progress_tone(0.8, tp) == "ahead", "比值 1.0 → 领先（边界含等号）"
    assert progress_tone(0.9, tp) == "ahead"
    assert progress_tone(0.56, tp) == "near", "比值 0.7 → 偏慢（边界含等号）"
    assert progress_tone(0.79, tp) == "near"
    assert progress_tone(0.55, tp) == "risk", "比值 < 0.7 → 滞后"
    assert progress_tone(0.1, 0) == "none", "时间进度为 0 时不做比值判断"


def test_status_is_consistent_across_dashboard_and_metrics(client, admin_headers):
    """回归：同一指标在驾驶舱与指标维护必须给出同一个状态。

    此前两处判据不同 —— 驾驶舱按「达成率 ÷ 时间进度」，指标维护按「实际 vs 目标」
    （年度目标未走完就一律红，年中看恒为「未达标」），同一指标会同时出现
    「领先」与「未达标」两种说法。
    """
    overview = client.get("/api/v1/dashboard/overview", params={"org_id": "CO"}, headers=admin_headers).json()
    board = client.get("/api/v1/metrics/board", params={"org_id": "CO"}, headers=admin_headers).json()

    ov = {r["metricId"]: r["tone"] for r in overview["annual"]}
    bd = {r["metricId"]: r["tone"] for r in board["annual"]}
    assert ov == bd, f"两处状态不一致：{ {k: (ov.get(k), bd.get(k)) for k in ov if ov.get(k) != bd.get(k)} }"

    # 项目公司单元格同样一致
    ov_cells = {
        (c["orgId"], m["metricId"]): m["tone"]
        for c in overview["projectCompare"]
        for m in c["cells"]
    }
    bd_cells = {
        (c["orgId"], m["metricId"]): m["tone"]
        for c in board["project"]["companies"]
        for m in c["cells"]
    }
    assert ov_cells == bd_cells, "项目公司单元格的状态在两处不一致"


def test_counts_follow_the_same_tone(client, admin_headers):
    """汇总统计（卡片右上角那组数字）必须与卡片标签同源。"""
    overview = client.get("/api/v1/dashboard/overview", params={"org_id": "CO"}, headers=admin_headers).json()
    counts = overview["counts"]

    assert set(counts) >= {"total", "done", "ahead", "near", "risk", "none", "filled"}
    assert counts["done"] + counts["ahead"] + counts["near"] + counts["risk"] + counts["none"] == counts["total"]

    # 与 annual 行的 tone 分布一致
    from collections import Counter

    dist = Counter(r["tone"] for r in overview["annual"])
    for key in ("done", "ahead", "near", "risk", "none"):
        assert counts[key] == dist.get(key, 0), f"{key} 的计数与卡片标签不一致"


def test_summary_uses_unified_wording(client, admin_headers):
    """分析结论文案用与卡片一致的词，不再出现「达标 / 未达标」这套旧说法。"""
    s = client.get("/api/v1/dashboard/summary", params={"org_id": "CO"}, headers=admin_headers).json()

    assert "verdict" in s, "整体状态字段应为 verdict（与单指标 tone 区分开）"
    assert s["verdict"] in ("good", "warn", "bad", "none")
    assert "已达成" in s["headline"] and "滞后" in s["headline"], s["headline"]
    assert "未达标" not in s["headline"], "不应再出现旧的「未达标」措辞"
    assert "领先" in s["headline"] and "偏慢" in s["headline"], s["headline"]


def test_summary_verdict_uses_the_same_rule_as_metrics(client, admin_headers):
    """整体结论必须与单指标同一套判定（达成率 ÷ 时间进度），阈值 ≥1 / 0.7。

    回归点：原先整体结论按绝对值判（加权达成率 < 0.9 → 存在风险），
    年中会出现自相矛盾的表述 —— 实测加权 83.8% 明明跑赢时间进度 78.4%，
    文案却写「整体低于时间轴」。
    """
    from app.services.achievement import progress_tone

    s = client.get("/api/v1/dashboard/summary", params={"org_id": "CO"}, headers=admin_headers).json()
    if s["weighted"] is None:
        assert s["verdict"] == "none"
        return

    # 与单指标同一函数、同一阈值
    expect = {"done": "good", "ahead": "good", "near": "warn", "risk": "bad", "none": "none"}[
        progress_tone(s["weighted"], s["timeProgress"])
    ]
    assert s["verdict"] == expect, f"加权 {s['weighted']:.3f} / 时间进度 {s['timeProgress']:.3f} → 应为 {expect}"

    # 具体的边界要求：跑赢时间进度就不该判为「存在风险」
    if s["weighted"] >= s["timeProgress"]:
        assert s["verdict"] in ("good",), f"跑赢时间进度却判为 {s['verdict']}"
    # 明显滞后（< 0.7 倍时间进度）才应判为 bad
    if s["weighted"] >= 0.7 * s["timeProgress"] and s["weighted"] < s["timeProgress"]:
        assert s["verdict"] == "warn", f"略慢于时间进度应判 warn，实际 {s['verdict']}"

    # 措辞与时间进度挂钩，不再有与事实相反的「低于时间轴」
    assert "时间进度" in s["headline"]


# ------------------------------ 部门任务完成度：展示口径筛选


def _dept_stats(client, headers, period: str, filter: str) -> list[dict]:
    return client.get(
        "/api/v1/dashboard/dept-task-stats",
        params={"period": period, "filter": filter},
        headers=headers,
    ).json()


def _stat_of(stats: list[dict], dept_id: str) -> dict | None:
    return next((s for s in stats if s["deptId"] == dept_id), None)


class _TempDept:
    """临时部门 + 指定创建时间的任务，用完即删（避免污染其他用例的统计口径）。"""

    def __init__(self, client, headers, specs: list[tuple[datetime, float]]):
        self.client = client
        self.headers = headers
        self.task_ids: list[str] = []
        self.dept_id = ""
        self._build(specs)

    def _build(self, specs):
        # 挂在「全公司」下才是部门（挂部门下会成为班组，统计会归集到上级部门）
        dept = self.client.post(
            "/api/v1/orgs",
            json={
                "name": f"临时口径部门{str(int(time.time() * 1000))[-6:]}",
                "parent_id": "CO",
                "head": "临时",
            },
            headers=self.headers,
        ).json()
        self.dept_id = dept["id"]

        for when, done in specs:
            body = _task_body(dept_id=self.dept_id)
            body["done"] = done
            r = self.client.post("/api/v1/tasks", json=body, headers=self.headers)
            assert r.status_code == 201, r.text
            task_id = r.json()["id"]
            self.task_ids.append(task_id)
            # created_at 由数据库默认写入，测试里显式回填以构造确定的创建时间
            with SessionLocal() as db:
                db.get(Task, task_id).created_at = when
                db.commit()

    def cleanup(self):
        for tid in self.task_ids:
            self.client.delete(f"/api/v1/tasks/{tid}", headers=self.headers)
        self.client.delete(f"/api/v1/orgs/{self.dept_id}", headers=self.headers)


def test_dept_task_stats_filters(client, admin_headers):
    """部门任务完成度支持 全部 / 本年新增 / 当月新增 / 未结束。

    口径：year / month 取「所选时间范围内创建的任务」，open 取未完成（done < target）。
    """
    now = datetime.now()
    month_start = datetime(now.year, now.month, 1)
    next_month = datetime(now.year + 1, 1, 1) if now.month == 12 else datetime(now.year, now.month + 1, 1)
    this_month = now
    this_year = datetime(now.year, 1, 15)
    last_year = datetime(now.year - 1, 6, 15)

    specs = [(this_month, 0), (this_year, 30), (last_year, 100)]
    temp = _TempDept(client, admin_headers, specs)
    try:
        period = now.strftime("%Y-%m")
        expect_month = sum(1 for when, _ in specs if month_start <= when < next_month)
        expect_open = sum(1 for _, done in specs if done < 100)

        all_stats = _dept_stats(client, admin_headers, period, "all")
        month_stats = _dept_stats(client, admin_headers, period, "month")
        year_stats = _dept_stats(client, admin_headers, period, "year")
        open_stats = _dept_stats(client, admin_headers, period, "open")

        assert _stat_of(all_stats, temp.dept_id)["total"] == 3, "全部口径应含 3 条"
        assert _stat_of(all_stats, temp.dept_id)["done"] == 1, "全部口径已完成的只有往年那条"
        assert _stat_of(month_stats, temp.dept_id)["total"] == expect_month, "当月新增只统计本月创建的任务"
        assert _stat_of(year_stats, temp.dept_id)["total"] == 2, "本年新增应排除往年创建的那条"
        assert _stat_of(open_stats, temp.dept_id)["total"] == expect_open, "未结束只统计未完成的任务"
        assert _stat_of(open_stats, temp.dept_id)["done"] == 0, "未结束口径下不应出现已完成"
        # 未识别的筛选值按「全部」处理，避免前端传错就白屏
        assert _dept_stats(client, admin_headers, period, "unknown") == all_stats
    finally:
        temp.cleanup()


def test_dept_task_stats_follows_selected_period(client, admin_headers):
    """年/月的取数跟着页面顶部周期走，而不是服务器当天。"""
    now = datetime.now()
    specs = [(now, 0), (datetime(now.year - 1, 6, 15), 100)]
    temp = _TempDept(client, admin_headers, specs)
    try:
        old_period = f"{now.year - 1}-06"
        year_stats = _dept_stats(client, admin_headers, old_period, "year")
        month_stats = _dept_stats(client, admin_headers, old_period, "month")

        assert _stat_of(year_stats, temp.dept_id)["total"] == 1, "切到往年周期后，本年新增只剩往年那条"
        assert _stat_of(month_stats, temp.dept_id)["total"] == 1, "当月新增应按所选月份取数"

        # 切到没有任务的月份 → 该部门不出现在结果里
        empty_month = f"{now.year - 1}-07"
        assert _stat_of(_dept_stats(client, admin_headers, empty_month, "month"), temp.dept_id) is None
    finally:
        temp.cleanup()


# ------------------------------ 项目管理（市场项目）与驾驶舱漏斗


def _mp_body(**over) -> dict:
    body = {
        "name": "变更测试项目",
        "description": "回归测试用",
        "stage": 25,
        "owner": "郑凯",
        "progress": "初始进展：已启动前期沟通",
    }
    body.update(over)
    return body


def test_market_project_crud_and_auto_fields(client, admin_headers):
    """新建自动生成 ID/创建人/时间戳；编辑更新字段；删除生效。"""
    r = client.post("/api/v1/market-projects", json=_mp_body(), headers=admin_headers)
    assert r.status_code == 201, r.text
    p = r.json()
    assert p["id"].startswith("MP")
    assert p["createdByName"] == "陈志明", "创建人应自动记录为当前登录人"
    assert p["createdAt"] and p["updatedAt"]
    assert p["stageLabel"] == "市场尽调"

    r = client.patch(
        f"/api/v1/market-projects/{p['id']}",
        json=_mp_body(
            name="变更测试项目-改",
            stage=90,
            tractorRequirements=[{"model": "6*4牵引车", "battery": "400", "qty": 66}],
            trailerRequirements=[{"model": "侧翻挂", "battery": "200", "qty": 33}],
        ),
        headers=admin_headers,
    )
    assert r.status_code == 200, r.text
    q = r.json()
    assert q["name"] == "变更测试项目-改" and q["stage"] == 90
    assert q["tractorDemand"] == 66 and q["trailerDemand"] == 33
    assert q["stageLabel"] == "签约落地"

    r = client.delete(f"/api/v1/market-projects/{p['id']}", headers=admin_headers)
    assert r.status_code == 200
    assert client.get(f"/api/v1/market-projects/{p['id']}", headers=admin_headers).status_code == 404


def test_market_project_stage_validation(client, admin_headers):
    """阶段必须是 10/25/50/75/90/100 之一；需求数不能为负。"""
    r = client.post("/api/v1/market-projects", json=_mp_body(stage=33), headers=admin_headers)
    assert r.status_code == 422
    r = client.post(
        "/api/v1/market-projects",
        json=_mp_body(tractorRequirements=[{"model": "x", "battery": "y", "qty": -1}]),
        headers=admin_headers,
    )
    assert r.status_code == 422


def test_market_project_scope_isolation(client, admin_headers, member_headers, head_headers):
    """数据范围按「创建人归属部门」匹配：self 只看本人创建；dept_sub 看本部门及下属创建；all 全量。"""
    # member（zhengkai, D01, self）创建一个项目
    r = client.post("/api/v1/market-projects", json=_mp_body(name="成员自建项目"), headers=member_headers)
    assert r.status_code == 201, r.text
    member_pid = r.json()["id"]
    # head（wangjg, D06, dept_sub）创建一个项目
    r = client.post("/api/v1/market-projects", json=_mp_body(name="负责人自建项目"), headers=head_headers)
    assert r.status_code == 201, r.text
    head_pid = r.json()["id"]
    # admin 创建一个（创建人属 CO 全公司）
    r = client.post("/api/v1/market-projects", json=_mp_body(name="管理员项目"), headers=admin_headers)
    assert r.status_code == 201, r.text
    admin_pid = r.json()["id"]

    try:
        member_ids = {p["id"] for p in client.get("/api/v1/market-projects", headers=member_headers).json()["items"]}
        assert member_pid in member_ids
        assert head_pid not in member_ids, "self 范围看不到别的部门创建的项目"
        assert admin_pid not in member_ids, "self 范围看不到全公司账号创建的项目"

        head_ids = {p["id"] for p in client.get("/api/v1/market-projects", headers=head_headers).json()["items"]}
        assert head_pid in head_ids
        assert member_pid not in head_ids, "dept_sub(D06) 看不到 D01 创建的项目"
        assert admin_pid not in head_ids, "dept_sub(D06) 看不到 CO 创建的项目"

        admin_ids = {p["id"] for p in client.get("/api/v1/market-projects", params={"size": 500}, headers=admin_headers).json()["items"]}
        assert {member_pid, head_pid, admin_pid} <= admin_ids, "all 范围应全量可见"

        # 范围外编辑/详情 → 403
        assert client.get(f"/api/v1/market-projects/{admin_pid}", headers=member_headers).status_code == 403
        r = client.patch(f"/api/v1/market-projects/{admin_pid}", json=_mp_body(), headers=member_headers)
        assert r.status_code == 403
        r = client.delete(f"/api/v1/market-projects/{member_pid}", headers=head_headers)
        assert r.status_code in (403, 404), "范围外删除同样被拒绝"
    finally:
        for pid, hdrs in ((member_pid, admin_headers), (head_pid, admin_headers), (admin_pid, admin_headers)):
            client.delete(f"/api/v1/market-projects/{pid}", headers=hdrs)


def test_market_menu_permission_required(client, admin_headers):
    """没有 market.view 的角色访问项目管理 → 403；但只要有 dashboard.view 就能看漏斗。"""
    r = client.post(
        "/api/v1/roles",
        json={"name": f"仅驾驶舱角色{str(int(time.time() * 1000))[-6:]}", "description": "测试", "scope": "all",
              "permissions": ["dashboard.view"]},
        headers=admin_headers,
    )
    assert r.status_code in (200, 201), r.text
    role_id = r.json()["id"]
    r = client.post(
        "/api/v1/users",
        json={"name": "漏斗观察员", "account": f"funnel{str(int(time.time() * 1000))[-6:]}", "password": "123456",
              "dept_id": "D01", "role_id": role_id, "scope": "all"},
        headers=admin_headers,
    )
    assert r.status_code in (200, 201), r.text
    uid = r.json()["id"]

    headers, _ = login(client, r.json()["account"])
    assert client.get("/api/v1/market-projects", headers=headers).status_code == 403
    assert client.get("/api/v1/dashboard/project-funnel", headers=headers).status_code == 200, "驾驶舱漏斗只看功能权限"
    assert client.get("/api/v1/dashboard/funnel-config", headers=headers).status_code == 200
    assert client.put("/api/v1/dashboard/funnel-config", json={"probs": []}, headers=headers).status_code == 403, \
        "没有 dashboard.config 不能改转化概率"

    client.delete(f"/api/v1/users/{uid}", headers=admin_headers) if client.get(f"/api/v1/users/{uid}", headers=admin_headers).status_code == 200 else None
    client.delete(f"/api/v1/roles/{role_id}", headers=admin_headers)


def test_project_funnel_per_stage_and_forecast(client, admin_headers):
    """漏斗按各阶段当前数量统计（stage == 该层），不向上累计；预测 = Σ 各层需求 × 该层概率。

    漏斗只统计「负责人属于市场部子树的活跃用户」的项目（避免脏 owner 污染统计），
    与列表口径（按创建人可见范围）不同，故这里校验漏斗自身一致性与预测公式，
    而不再与列表逐条对账。
    """
    funnel = client.get("/api/v1/dashboard/project-funnel", headers=admin_headers).json()
    assert [s["stage"] for s in funnel["stages"]] == [10, 25, 50, 75, 90, 100]

    # 各层都是「停在该阶段」的当前快照（非负整数），且至少有种子项目被纳入，避免空漏斗误判
    total = 0
    for s in funnel["stages"]:
        assert s["count"] >= 0 and s["tractor"] >= 0 and s["trailer"] >= 0
        assert s["prob"] == {p["stage"]: p["prob"] for p in funnel["stages"]}[s["stage"]]
        total += s["count"]
    assert total > 0, "种子数据应至少有项目被漏斗纳入"

    # 预测 = Σ（各层需求 × 各层概率），与漏斗自身口径一致
    exp_trac = round(sum(s["tractor"] * s["prob"] for s in funnel["stages"]))
    exp_trai = round(sum(s["trailer"] * s["prob"] for s in funnel["stages"]))
    assert funnel["forecast"]["tractor"] == exp_trac
    assert funnel["forecast"]["trailer"] == exp_trai


def test_funnel_config_affects_forecast(client, admin_headers):
    """转化概率整份保存、立即影响预测；非法配置被拒绝；测完还原。"""
    cfg = client.get("/api/v1/dashboard/funnel-config", headers=admin_headers).json()["probs"]
    original = [{"stage": c["stage"], "prob": c["prob"]} for c in cfg]
    try:
        # 概率全 1.0 → 预测 = 各层需求之和（不再按阶段加权打折）
        all_one = [{"stage": c["stage"], "prob": 1.0} for c in cfg]
        r = client.put("/api/v1/dashboard/funnel-config", json={"probs": all_one}, headers=admin_headers)
        assert r.status_code == 200, r.text

        funnel = client.get("/api/v1/dashboard/project-funnel", headers=admin_headers).json()
        assert funnel["forecast"]["tractor"] == sum(s["tractor"] for s in funnel["stages"])
        assert funnel["forecast"]["trailer"] == sum(s["trailer"] for s in funnel["stages"])

        # 概率全 0.5 → 预测 = 各层需求 × 0.5 取整
        half = [{"stage": c["stage"], "prob": 0.5} for c in cfg]
        r = client.put("/api/v1/dashboard/funnel-config", json={"probs": half}, headers=admin_headers)
        assert r.status_code == 200, r.text
        funnel = client.get("/api/v1/dashboard/project-funnel", headers=admin_headers).json()
        assert funnel["forecast"]["tractor"] == round(sum(s["tractor"] * 0.5 for s in funnel["stages"]))
        assert funnel["forecast"]["trailer"] == round(sum(s["trailer"] * 0.5 for s in funnel["stages"]))

        # 缺阶段 / 概率越界 → 422
        assert client.put("/api/v1/dashboard/funnel-config", json={"probs": all_one[:-1]}, headers=admin_headers).status_code == 422
        bad = [dict(p) for p in all_one]
        bad[0]["prob"] = 1.5
        assert client.put("/api/v1/dashboard/funnel-config", json={"probs": bad}, headers=admin_headers).status_code == 422
    finally:
        client.put("/api/v1/dashboard/funnel-config", json={"probs": original}, headers=admin_headers)
